from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

ROOT = Path(r"F:\PUN\Tesis-Mundo-Aprendo\Mundo aprendo tesis")
OUT = ROOT / "Informacion_tecnica_Mundo_aprendo_tesis.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(90, 90, 90)
LIGHT = "E8EEF5"
WHITE = RGBColor(255, 255, 255)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.2)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 14, 7),
    ("Heading 2", 13, BLUE, 10, 5),
    ("Heading 3", 11.5, NAVY, 7, 3),
]:
    s = styles[name]
    s.font.name = "Calibri"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = color
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = t.rows[0].cells[i]
        c.width = Inches(w)
        shade(c, LIGHT)
        margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.2)
        r.font.color.rgb = NAVY
    set_repeat_table_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, (value, w) in enumerate(zip(row, widths)):
            cells[i].width = Inches(w)
            margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.38)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)

header = sec.header.paragraphs[0]
header.text = "MUNDO APRENDO TESIS  |  INFORMACIÓN TÉCNICA VERIFICADA"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.color.rgb = GRAY

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Documento elaborado a partir de la configuración, código, compilaciones y pruebas presentes en el proyecto.")
r.font.size = Pt(8)
r.font.color.rgb = GRAY

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(9)
p.paragraph_format.space_after = Pt(3)
r = p.add_run("INFORMACIÓN TÉCNICA DEL PROYECTO")
r.font.name = "Calibri"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = NAVY
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(15)
r = p2.add_run("Mundo aprendo tesis")
r.font.size = Pt(15)
r.font.color.rgb = GRAY
r.italic = True

doc.add_heading("1. Identificación y compilación", level=1)
table(["Dato", "Información verificada"], [
    ("Nombre exacto del producto", "Mundo aprendo tesis"),
    ("Versión de Unity", "Unity 6000.4.0f1, revisión 8cf496087c8f"),
    ("Plataforma", "Windows Standalone"),
    ("Arquitectura", "x86-64 (PE machine 0x8664); el build contiene Plugins/x86_64 y UnityCrashHandler64.exe"),
    ("Backend observado en el build", "Mono: el paquete compilado contiene MonoBleedingEdge y ensamblados administrados"),
    ("Estado de compilación", "Correcto. UnityBuild-Windows-final.log registra “Build Finished, Result: Success” y retorno 0."),
], [1.65, 5.2])

doc.add_heading("2. Pantalla y requisitos de ejecución", level=1)
table(["Elemento", "Configuración o requisito"], [
    ("Resolución configurada", "1024 × 768 píxeles."),
    ("Modo de pantalla", "fullscreenMode = 1; ventana no redimensionable (resizableWindow = 0)."),
    ("Resoluciones verificadas por prueba", "1920 × 1080, 1366 × 768 y 1280 × 720; la prueba de controles dentro del Canvas figura aprobada."),
    ("Sistema operativo", "Windows de 64 bits. Es requisito para el build existente y para el servicio de dictado usado por el juego."),
    ("Entrada", "Mouse para la interfaz de botones; el proyecto también incluye Unity Input System 1.19.0."),
    ("Micrófono", "Necesario para MundoCuentos_VozTest y su reconocimiento de lectura. El menú de audio enumera, selecciona y prueba dispositivos."),
    ("Audio", "Salida de audio recomendada/necesaria para música, efectos y secuencias del Mundo Musical."),
], [1.7, 5.15])

doc.add_heading("3. Escenas y orden de navegación", level=1)
doc.add_paragraph("Orden registrado en Build Settings y flujo comprobado por los scripts de navegación:")
table(["Orden", "Escena", "Función en el flujo"], [
    ("1", "Menu", "Escena inicial; acceso a la selección de mundos y opciones."),
    ("2", "SeleccionMundos", "Selector central. Presenta estrellas, bloqueos y acceso a cada mundo."),
    ("3", "MundoMusical", "Minijuego de secuencias musicales."),
    ("4", "MundoCuentos_VozTest", "Lectura de cuentos con micrófono, dictado y evaluación."),
    ("5", "MundoTamanos", "Minijuego de comparación de tamaños con animales."),
    ("6", "MundoEmociones", "Minijuego de identificación de emociones."),
], [0.55, 2.05, 4.25])
doc.add_paragraph("Navegación: Menu → SeleccionMundos → mundo elegido. Desde los mundos se retorna a SeleccionMundos o Menu mediante SceneNavigation. El desbloqueo es lineal: el primer mundo está disponible y cada mundo posterior requiere completar el anterior.")

doc.add_heading("4. Scripts principales", level=1)
scripts = [
    ("Assets/Mundo Aprendo/Scripts/Navigation/MundoAprendoSceneNames.cs", "Centraliza los nombres exactos de las seis escenas."),
    ("Assets/Mundo Aprendo/Scripts/Navigation/SceneNavigation.cs", "Valida que la escena esté habilitada, ejecuta transición UI y carga escenas."),
    ("Assets/Mundo Aprendo/Scripts/WorldSelectionManager.cs", "Gestiona tarjetas de mundos, bloqueos, estrellas, selección y reinicio de progreso."),
    ("Assets/Mundo Aprendo/Scripts/Progress/WorldProgressRepository.cs", "Lee y guarda completado/estrellas por mundo; aplica desbloqueo lineal."),
    ("Assets/Mundo Aprendo/Scripts/Progress/StarRatingCalculator.cs", "Convierte cantidad de errores en 0–3 estrellas según umbrales."),
    ("Assets/Mundo Aprendo/Scripts/UI/UIStarDisplay.cs", "Actualiza y anima las imágenes de estrellas existentes."),
    ("Assets/Mundo Aprendo/Scripts/UI/UISceneTransition.cs", "Presenta transiciones visuales al cambiar de escena."),
    ("Assets/Mundo Aprendo/Scripts/AudioManager.cs", "Administra música/sonidos del menú y evita duplicados."),
    ("Assets/Mundo Aprendo/Scripts/Options/AudioOptionsController.cs", "Controla volumen, lista de micrófonos y prueba de señal/frecuencia."),
    ("Assets/Mundo Aprendo/Scripts/Options/MicrophoneSettings.cs", "Persiste el nombre del micrófono seleccionado y valida disponibilidad."),
    ("Assets/Mundo Aprendo/Scripts/MundoCuentos/Scripts/VoiceRecognitionTest.cs", "Coordina selección de cuento, micrófono, dictado, evaluación, resultado, estrellas y navegación."),
    ("Assets/Mundo Aprendo/Scripts/MundoCuentos/Scripts/WindowsDictationSpeechService.cs", "Encapsula DictationRecognizer de Windows y emite texto parcial/final y errores."),
    ("Assets/Mundo Aprendo/Scripts/MundoCuentos/Scripts/ReadingEvaluator.cs", "Normaliza y compara lectura reconocida con el cuento; calcula similitud y estrellas."),
    ("Assets/Mundo Aprendo/Scripts/MundoCuentos/Scripts/StoryProgressRepository.cs", "Guarda mejor puntaje, estrellas, estado completado y tutorial por cuento."),
    ("Assets/Mundo Aprendo/Scripts/MundoMusical/Scripts/MundoMusicalSequenceGame.cs", "Ejecuta secuencias musicales, controla entrada, errores, estrellas y finalización."),
    ("Assets/Mundo Aprendo/Scripts/SizeWorldController.cs", "Ejecuta rondas del minijuego de tamaños, respuestas y resultado."),
    ("Assets/Mundo Aprendo/Scripts/MundoEmociones/EmotionGameManager.cs", "Gestiona rondas, respuestas y progreso del minijuego de emociones."),
]
table(["Ruta", "Función"], scripts, [3.75, 3.1])

doc.add_heading("5. Almacenamiento del progreso", level=1)
doc.add_paragraph("El progreso se almacena con PlayerPrefs y se fuerza su escritura mediante PlayerPrefs.Save(). No se encontró un sistema de guardado en archivo propio, base de datos ni servicio remoto.")
bullet("Por mundo: claves MundoAprendo_World_{índice}_Completed y MundoAprendo_World_{índice}_Stars.")
bullet("Se conserva el mejor resultado de estrellas, limitado entre 0 y 3.")
bullet("Cuentos: StoryProgressRepository guarda completado, mejor puntaje y mejores estrellas por identificador estable.")
bullet("También se guardan estados de tutorial y el nombre del micrófono seleccionado (selected_microphone_device).")
bullet("WorldProgressRepository.ResetAll elimina el progreso de mundos, cuentos y tutoriales asociados.")

doc.add_heading("6. Reconocimiento de voz, paquetes y librerías", level=1)
table(["Componente", "Uso verificado"], [
    ("UnityEngine.Windows.Speech.DictationRecognizer", "Motor de reconocimiento de voz de Windows usado por WindowsDictationSpeechService."),
    ("UnityEngine.Microphone", "Detección de dispositivos, prueba de nivel/señal y validación previa al dictado."),
    ("ISpeechToTextService", "Interfaz interna que desacopla el controlador del servicio de dictado."),
    ("Unity Input System 1.19.0", "Paquete de entrada configurado en el proyecto."),
    ("Universal Render Pipeline 17.4.0", "Pipeline gráfico instalado y usado por el proyecto."),
    ("Unity Test Framework 1.6.0", "Infraestructura de pruebas EditMode y PlayMode."),
    ("TextMesh Pro / UGUI 2.0.0", "Texto e interfaz de usuario."),
], [2.55, 4.3])
doc.add_paragraph("No se encontró un SDK externo de voz, servicio web de transcripción ni librería de terceros para reconocimiento. La implementación depende de la API de dictado incluida en Unity para Windows.")

doc.add_heading("7. Módulos o sistemas principales", level=1)
table(["Sistema", "Comportamiento verificado"], [
    ("Navegación", "Nombres centralizados, validación de escenas habilitadas, transición visual y retorno a menú/selector."),
    ("Progreso", "PlayerPrefs por mundo y cuento; notificación ProgressChanged para refrescar la interfaz."),
    ("Estrellas", "Calificación de 0–3, conservación del mejor resultado y representación/animación visual."),
    ("Desbloqueo", "Secuencial entre cuatro mundos; además, Cuentos controla acceso según cuentos completados."),
    ("Minijuegos", "Musical (secuencias), Cuentos (lectura y voz), Tamaños (comparación) y Emociones (selección de emoción)."),
    ("Audio", "Música/efectos, control de volumen, selección/prueba de micrófono y reproducción de secuencias musicales."),
], [1.45, 5.4])

doc.add_heading("8. Pruebas funcionales y resultados", level=1)
table(["Prueba o conjunto", "Resultado verificado"], [
    ("EditMode final", "33 aprobadas, 0 fallidas, 0 omitidas. Registro: UnityTest-EditModeApi-final.log."),
    ("PlayMode", "8 aprobadas, 0 fallidas, 0 omitidas. XML: Logs/MundoCuentosExpansionPlayModeApi.xml."),
    ("Carga de escenas", "Aprobada: todas las escenas configuradas cargan con su UI principal."),
    ("Diseño responsivo", "Aprobada en 1920×1080, 1366×768 y 1280×720: controles activos dentro del Canvas."),
    ("Persistencia", "Aprobada: conserva mejores estrellas; usa IDs estables para cuentos y exige cuentos distintos."),
    ("Audio de menú", "Aprobada: se limita al menú y no se duplica."),
    ("Transiciones UI", "Aprobada: panel abre/cierra y restaura interacción."),
    ("Estrellas", "Aprobada: actualiza tres imágenes persistentes."),
    ("Compilación Windows", "Aprobada: Build Finished, Result: Success; retorno del proceso 0."),
], [2.15, 4.7])
doc.add_paragraph("Nota de trazabilidad: existe un XML EditMode más antiguo con 44 aprobadas y 5 fallidas por rutas obsoletas Assets/Scenes/...; el registro final posterior informa 33/0. Se conserva como evidencia histórica de una limitación de las pruebas anteriores, no como resultado final vigente.")

doc.add_heading("9. Errores conocidos, limitaciones y pendientes", level=1)
bullet("DictationRecognizer usa el micrófono predeterminado de Windows. Aunque el juego permite elegir y probar un dispositivo con UnityEngine.Microphone, Unity no permite pasar directamente esa selección al DictationRecognizer.")
bullet("El reconocimiento de voz implementado es específico de Windows; no hay una implementación alternativa confirmada para otras plataformas.")
bullet("El uso del micrófono requiere que Windows tenga un dispositivo disponible y permisos/servicios de voz operativos; el código muestra mensajes de error cuando no hay señal o dispositivo.")
bullet("En la selección de cuentos, los cuentos configurados como no disponibles se muestran como “Muy pronto”; por tanto, no todo el contenido listado necesariamente está habilitado.")
bullet("Permanece un resultado EditMode histórico fallido por rutas antiguas de escenas. El resultado final posterior es correcto, pero conviene evitar usar ese XML antiguo como reporte vigente.")

doc.add_heading("Fuentes verificadas dentro del proyecto", level=1)
for src in [
    "ProjectSettings/ProjectVersion.txt",
    "ProjectSettings/ProjectSettings.asset",
    "ProjectSettings/EditorBuildSettings.asset",
    "Packages/manifest.json",
    "Assets/Mundo Aprendo/Scripts/**",
    "Assets/Editor/Tests/** y Assets/Tests/PlayMode/**",
    "UnityBuild-Windows-final.log",
    "UnityTest-EditModeApi-final.log",
    "Logs/MundoCuentosExpansionPlayModeApi.xml",
    "Build/ y Builds/Windows/",
]:
    bullet(src)

doc.core_properties.title = "Información técnica - Mundo aprendo tesis"
doc.core_properties.subject = "Información verificada del proyecto Unity"
doc.core_properties.author = "Codex"
doc.save(OUT)
print(OUT)
